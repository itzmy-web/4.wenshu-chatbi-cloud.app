# 数据导入路由：POST /api/upload 上传 CSV/Excel → 动态建表 → LLM 自动生成字段语义 → 立即可问数
import csv
import io
import json
import re
import uuid

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.params import Depends as DependsP
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import PromptTemplate
from sqlalchemy import text

from app.agent.llm import llm
from app.api.dependencies import get_dw_session, get_meta_session
from app.api.routers.auth_router import verify_token
from app.clients.embedding_client_manager import embedding_client_manager
from app.core.log import logger
from app.models.mysql.column_info_mysql import ColumnInfoMySQL
from app.models.mysql.table_info_mysql import TableInfoMySQL

upload_router = APIRouter()

MAX_ROWS = 20000
MAX_COLS = 100
NUM_RE = re.compile(r"^-?\d+(\.\d+)?$")


def _parse_csv(content: bytes) -> list[list[str]]:
    for enc in ("utf-8-sig", "gbk", "utf-16"):
        try:
            text_csv = content.decode(enc)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    else:
        raise HTTPException(status_code=400, detail="文件编码无法识别，请另存为 UTF-8 的 CSV")
    rows = list(csv.reader(io.StringIO(text_csv)))
    return [r for r in rows if any(cell.strip() for cell in r)]


def _parse_xlsx(content: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(status_code=400, detail="服务端缺少 Excel 解析组件（openpyxl）")
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else str(v) for v in row])
    wb.close()
    return [r for r in rows if any(cell.strip() for cell in r)]


def _slug(name: str) -> str:
    s = re.sub(r"[^0-9a-zA-Z\u4e00-\u9fa5]+", "_", name).strip("_").lower()
    return s[:24] or "data"


async def _llm_table_meta(table: str, columns: list[dict], sample: list[dict]) -> dict:
    """让 LLM 根据表名/列名/样本行生成中文表描述与字段描述、别名。失败时返回兜底结果。"""
    fallback = {"description": f"用户上传的数据表 {table}",
                "columns": [{"name": c["name"], "description": c["name"], "alias": []} for c in columns]}
    prompt = PromptTemplate(
        template=(
            "你是数据仓库专家。下面是一张数据表的表名、字段和样本行，请推断每个字段的业务含义。\n"
            "表名：{table}\n字段：{columns}\n样本行：{sample}\n\n"
            "只输出 JSON（不要 markdown 代码块），格式：\n"
            '{{"description": "一句话中文表描述", "columns": [{{"name": "字段名", "description": "中文含义", '
            '"alias": ["中文别名1", "别名2"]}}]}}\n'
            "要求：description/alias 用简短中文；alias 给 1~3 个用户常用的叫法；字段名必须与给定字段一一对应。"
        ),
        input_variables=["table", "columns", "sample"],
    )
    try:
        chain = prompt | llm | StrOutputParser()
        raw = await chain.ainvoke({
            "table": table,
            "columns": json.dumps([c["name"] for c in columns], ensure_ascii=False),
            "sample": json.dumps(sample[:5], ensure_ascii=False, default=str)[:1200],
        })
        data = json.loads(re.sub(r"^```(json)?|```$", "", raw.strip(), flags=re.MULTILINE).strip())
        col_map = {c["name"]: c for c in data.get("columns", []) if isinstance(c, dict) and "name" in c}
        out_cols = []
        for c in columns:
            m = col_map.get(c["name"], {})
            out_cols.append({"name": c["name"],
                             "description": str(m.get("description") or c["name"]),
                             "alias": [str(a) for a in (m.get("alias") or [])][:3]})
        return {"description": str(data.get("description") or f"用户上传的数据表 {table}"), "columns": out_cols}
    except Exception as e:
        logger.warning(f"LLM 生成表语义失败（使用兜底）：{e}")
        return fallback


@upload_router.post("/api/upload")
async def upload_data(file: UploadFile = File(...),
                      _token: None = DependsP(verify_token),
                      meta_session=DependsP(get_meta_session),
                      dw_session=DependsP(get_dw_session)):
    filename = file.filename or "data.csv"
    content = await file.read()
    if len(content) > 20 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件过大（限 20MB）")
    lower = filename.lower()
    if lower.endswith(".csv"):
        raw_rows = _parse_csv(content)
    elif lower.endswith(".xlsx"):
        raw_rows = _parse_xlsx(content)
    else:
        raise HTTPException(status_code=400, detail="仅支持 .csv 或 .xlsx 文件")
    if len(raw_rows) < 2:
        raise HTTPException(status_code=400, detail="文件至少需要表头和一行数据")

    header = [h.strip() or f"col{i}" for i, h in enumerate(raw_rows[0][:MAX_COLS])]
    data_rows = []
    for r in raw_rows[1: MAX_ROWS + 1]:
        r = list(r[:MAX_COLS]) + [""] * (len(header) - len(r[:MAX_COLS]))
        data_rows.append(dict(zip(header, r)))
    if not data_rows:
        raise HTTPException(status_code=400, detail="没有解析到数据行")

    # 表名：u_ + 文件名 slug + 4位随机，避免冲突
    ext_re = re.compile(r"\.(csv|xlsx)$")
    stem = _slug(ext_re.sub("", filename.lower()))
    table = f"u_{stem}_{uuid.uuid4().hex[:4]}"

    # 类型推断：全为数字的列 → number，否则 text
    col_types = {}
    for c in header:
        values = [r[c] for r in data_rows if str(r.get(c, "")).strip() != ""]
        col_types[c] = "number" if values and all(NUM_RE.match(str(v)) for v in values) else "text"

    # 1) dw.db 动态建表 + 写入数据
    cols_ddl = ", ".join(f'"{c}" {"REAL" if col_types[c] == "number" else "TEXT"}' for c in header)
    await dw_session.execute(text(f'CREATE TABLE "{table}" ({cols_ddl})'))
    insert_sql = f'INSERT INTO "{table}" ({", ".join(chr(34) + c + chr(34) for c in header)}) ' \
                 f'VALUES ({", ".join(":" + c for c in header)})'
    params = [{c: (float(r[c]) if col_types[c] == "number" and str(r[c]).strip() != "" else str(r.get(c, "")))
               for c in header} for r in data_rows]
    for i in range(0, len(params), 500):
        await dw_session.execute(text(insert_sql), params[i:i + 500])
    await dw_session.commit()

    # 2) LLM 生成表/字段中文语义
    meta_gen = await _llm_table_meta(table, [{"name": c, "type": col_types[c]} for c in header], data_rows[:5])

    # 3) 写 meta：表 + 字段
    tid = f"tbl_{uuid.uuid4().hex[:12]}"
    meta_session.add(TableInfoMySQL(id=tid, name=table, role="fact", description=meta_gen["description"]))
    embedder = embedding_client_manager.client
    for c in header:
        m = next(x for x in meta_gen["columns"] if x["name"] == c)
        cid = f"{table}.{c}"
        examples = []
        for r in data_rows:
            v = str(r[c])
            if v and v not in examples:
                examples.append(v)
            if len(examples) >= 5:
                break
        meta_session.add(ColumnInfoMySQL(
            id=cid, name=c, type=col_types[c],
            role="measure" if col_types[c] == "number" else "dimension",
            examples=examples, description=m["description"], alias=m["alias"], table_id=tid))
        # 4) 词汇向量（列名 + 描述 + 别名）
        payload = {"id": cid, "name": c, "type": col_types[c],
                   "role": "measure" if col_types[c] == "number" else "dimension",
                   "examples": examples, "description": m["description"], "alias": m["alias"], "table_id": tid}
        for t in [c, m["description"]] + m["alias"]:
            if t:
                vec = await embedder.aembed_query(t)
                await meta_session.execute(text(
                    "INSERT OR REPLACE INTO vectors(collection, point_id, payload_json, vector_json) "
                    "VALUES ('column', :p, :pj, :vj)"),
                    {"p": uuid.uuid4().hex, "pj": json.dumps(payload, ensure_ascii=False),
                     "vj": json.dumps(vec)})
        # 5) 低基数文本列的取值进 value_info（供取值召回）
        if col_types[c] == "text":
            distinct = {str(r[c]) for r in data_rows if str(r[c]).strip()} 
            if 0 < len(distinct) <= 50:
                for v in distinct:
                    await meta_session.execute(text(
                        "INSERT OR REPLACE INTO value_info(id, value, type, column_id, column_name, table_id, table_name) "
                        "VALUES (:id, :value, :type, :column_id, :column_name, :table_id, :table_name)"),
                        {"id": uuid.uuid4().hex, "value": v, "type": "dimension", "column_id": cid,
                         "column_name": c, "table_id": tid, "table_name": table})
    await meta_session.commit()

    logger.info(f"上传入库成功：{filename} -> {table}（{len(data_rows)} 行 {len(header)} 列）")
    return {"table": table, "rows": len(data_rows),
            "columns": [{"name": c, "type": col_types[c],
                         "description": next(x for x in meta_gen["columns"] if x["name"] == c)["description"],
                         "alias": next(x for x in meta_gen["columns"] if x["name"] == c)["alias"]} for c in header]}
